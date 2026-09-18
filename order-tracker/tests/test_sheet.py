from pathlib import Path

from openpyxl import Workbook

from app.sheet import apply_results, load_table


def _sample(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["送货日期", "承运商", "Tracking No.", "状态", "Delivery Address", "Notes"])
    ws.append(["", "USPS", "9214490422577138214308", "签收", "Anderson Werner\nUT", ""])
    ws.append(["", "8DT", "EWSMM260824000268YQ", "", "Axel Medina\nMexico", ""])
    ws.append(["", "UPS", "1Z1B0D750301336987", "", "Brenda Vincent\nLA", "海运"])
    ws.append(["20260909", "", "", "", "Archer Rosenkrantz\nDenver", ""])
    wb.save(path)
    return path


def test_load_and_filter(tmp_path: Path):
    table = load_table(_sample(tmp_path / "s.xlsx"))
    assert table.columns.tracking == 3
    assert table.columns.carrier == 2
    assert table.delivered_count == 1
    assert table.pending_count == 2
    assert table.missing_count == 1
    pending = [row for row in table.rows if row.needs_query]
    assert {row.recipient for row in pending} == {"Axel Medina", "Brenda Vincent"}
    missing = [row for row in table.rows if row.skip_reason == "missing_tracking"]
    assert missing[0].recipient == "Archer Rosenkrantz"


def test_export_template_missing_tracking(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["参考号", "运单号", "渠道", "收件人", "收件地址"])
    ws.append(["YT091701", None, None, None, "Diego Cea\nFillmore, CA"])
    ws.append(["YT091705", None, None, "Elizabeth", "Hilliard, OH"])
    path = tmp_path / "export.xlsx"
    wb.save(path)
    table = load_table(path)
    assert table.columns.tracking == 2
    assert table.columns.carrier == 3
    assert table.pending_count == 0
    assert table.missing_count == 2
    assert {row.recipient for row in table.rows} == {"Diego Cea", "Elizabeth"}


def test_explicit_columns_and_sheet(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "新模板"
    ws.append(["参考号", "运单号", "渠道", "收件人"])
    ws.append(["YT1", "EWSMM260824000268YQ", "8DT", "Axel"])
    other = wb.create_sheet("Sheet1")
    other.append(["A", "B"])
    other.append(["x", "y"])
    path = tmp_path / "multi.xlsx"
    wb.save(path)
    table = load_table(path, sheet_name="新模板", tracking_col=2, carrier_col=3)
    assert table.sheet_name == "新模板"
    assert table.columns.tracking == 2
    assert table.columns.carrier == 3
    assert table.pending_count == 1
    assert table.rows[0].carrier_raw == "8DT"
    assert table.rows[0].tracking_numbers == ["EWSMM260824000268YQ"]


def test_apply_results_writes_status(tmp_path: Path):
    source = _sample(tmp_path / "s.xlsx")
    dest = tmp_path / "out.xlsx"
    apply_results(
        source,
        dest,
        {
            3: {"status": "签收", "latest": "墨西哥当地签收", "queried_at": "2026-09-18 15:00"},
            5: {"status": "未填单号", "latest": "", "queried_at": "2026-09-18 15:00"},
        },
    )
    table = load_table(dest)
    by_row = {row.excel_row: row for row in table.rows}
    assert by_row[3].status_raw == "签收"
    assert by_row[5].status_raw == "未填单号"
    assert by_row[2].status_raw == "签收"
