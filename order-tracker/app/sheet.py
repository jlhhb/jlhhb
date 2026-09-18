"""解析发货清单 xlsx：自动识别承运商/运单号/状态列，并写回查询结果。"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .status import is_delivered_text

TRACKING_TOKEN_RE = re.compile(
    r"(?:"
    r"EWS[A-Z0-9]{10,}YQ|"
    r"1Z[A-Z0-9]{16}|"
    r"9[234]\d{18,24}|"
    r"8412\d{8,14}|"
    r"81600\d{6,12}|"
    r"87\d{10,14}|"
    r"\d{16}"
    r")",
    re.I,
)

CARRIER_HINTS = {
    "usps": "usps",
    "ups": "ups",
    "fedex": "fedex",
    "canada post": "canada_post",
    "canadapost": "canada_post",
    "dpd": "dpd",
    "8dt": "8dt",
    "永利": "8dt",
    "tgx": "tgx",
}


@dataclass
class SheetRow:
    excel_row: int
    recipient: str
    carrier_raw: str
    tracking_raw: str
    tracking_numbers: list[str]
    status_raw: str
    address: str
    notes: str
    skip_reason: str | None = None

    @property
    def needs_query(self) -> bool:
        if self.skip_reason:
            return False
        if is_delivered_text(self.status_raw):
            return False
        return bool(self.tracking_numbers)


@dataclass
class ColumnMap:
    carrier: int | None
    tracking: int
    status: int
    address: int | None
    notes: int | None
    latest: int
    queried_at: int
    recipient_from: int | None


@dataclass
class SheetTable:
    path: Path
    headers: list[str]
    columns: ColumnMap
    rows: list[SheetRow]
    delivered_count: int
    pending_count: int
    missing_count: int


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_name(value: Any, index: int) -> str:
    text = _cell_text(value)
    return text if text else f"列{index}"


def extract_tracking_numbers(text: str) -> list[str]:
    found = TRACKING_TOKEN_RE.findall(text or "")
    seen: list[str] = []
    for token in found:
        item = token.strip()
        if item and item not in seen:
            seen.append(item)
    return seen


def _score_tracking_column(ws: Worksheet, col: int) -> int:
    score = 0
    for row in range(2, min(ws.max_row + 1, 80)):
        score += len(extract_tracking_numbers(_cell_text(ws.cell(row, col).value)))
    return score


def _looks_like_carrier(text: str) -> bool:
    lowered = text.lower()
    return any(key in lowered for key in CARRIER_HINTS)


def _looks_like_status(text: str) -> bool:
    return is_delivered_text(text) or text in {"运输中", "已上网", "仅面单", "官网查无", "未填单号"}


def detect_columns(ws: Worksheet) -> ColumnMap:
    max_col = ws.max_column or 1
    headers = [_cell_text(ws.cell(1, c).value).lower() for c in range(1, max_col + 1)]

    tracking_col = None
    for idx, header in enumerate(headers, start=1):
        if any(key in header for key in ("tracking", "运单", "快递单", "物流单")):
            tracking_col = idx
            break
    if tracking_col is None:
        scored = [(c, _score_tracking_column(ws, c)) for c in range(1, max_col + 1)]
        scored.sort(key=lambda item: item[1], reverse=True)
        if scored and scored[0][1] > 0:
            tracking_col = scored[0][0]
    if tracking_col is None:
        raise ValueError("没有识别到运单号列，请确认表格中有 Tracking No. 或运单号")

    carrier_col = None
    for idx, header in enumerate(headers, start=1):
        if any(key in header for key in ("carrier", "承运", "物流公司", "快递公司")):
            carrier_col = idx
            break
    if carrier_col is None and tracking_col > 1:
        hint_hits = 0
        probe = tracking_col - 1
        for row in range(2, min(ws.max_row + 1, 40)):
            if _looks_like_carrier(_cell_text(ws.cell(row, probe).value)):
                hint_hits += 1
        if hint_hits:
            carrier_col = probe

    status_col = None
    for idx, header in enumerate(headers, start=1):
        if any(key in header for key in ("status", "状态", "签收")):
            status_col = idx
            break
    if status_col is None:
        probe = tracking_col + 1
        if probe <= max_col:
            hits = 0
            for row in range(2, min(ws.max_row + 1, 40)):
                if _looks_like_status(_cell_text(ws.cell(row, probe).value)):
                    hits += 1
            if hits:
                status_col = probe
        if status_col is None:
            status_col = min(tracking_col + 1, max(max_col, tracking_col + 1))
            if status_col > max_col:
                status_col = tracking_col + 1

    address_col = None
    notes_col = None
    latest_col = None
    queried_col = None
    recipient_col = None
    for idx, header in enumerate(headers, start=1):
        if recipient_col is None and any(key in header for key in ("收件人", "recipient", "consignee")):
            recipient_col = idx
        if address_col is None and any(key in header for key in ("address", "地址")) and "收件人" not in header:
            address_col = idx
        if notes_col is None and any(key in header for key in ("note", "备注", "notes")):
            notes_col = idx
        if latest_col is None and any(key in header for key in ("最新轨迹", "latest", "轨迹")):
            latest_col = idx
        if queried_col is None and any(key in header for key in ("查询时间", "queried")):
            queried_col = idx

    latest_col = latest_col or (max(max_col, status_col, tracking_col) + 1)
    queried_col = queried_col or (latest_col + 1)

    return ColumnMap(
        carrier=carrier_col,
        tracking=tracking_col,
        status=status_col,
        address=address_col,
        notes=notes_col,
        latest=latest_col,
        queried_at=queried_col,
        recipient_from=recipient_col or address_col,
    )


def _recipient_from(address: str) -> str:
    for line in (address or "").splitlines():
        name = line.strip().strip('"')
        if name:
            return name
    return ""


def load_table(path: Path) -> SheetTable:
    wb = load_workbook(path)
    ws = wb.active
    columns = detect_columns(ws)
    headers = [_header_name(ws.cell(1, c).value, c) for c in range(1, (ws.max_column or 1) + 1)]
    rows: list[SheetRow] = []
    delivered = pending = missing = 0
    for excel_row in range(2, (ws.max_row or 1) + 1):
        tracking_raw = _cell_text(ws.cell(excel_row, columns.tracking).value)
        carrier_raw = _cell_text(ws.cell(excel_row, columns.carrier).value) if columns.carrier else ""
        status_raw = _cell_text(ws.cell(excel_row, columns.status).value)
        address = _cell_text(ws.cell(excel_row, columns.address).value) if columns.address else ""
        notes = _cell_text(ws.cell(excel_row, columns.notes).value) if columns.notes else ""
        recipient = ""
        if columns.recipient_from:
            recipient = _recipient_from(_cell_text(ws.cell(excel_row, columns.recipient_from).value))
        if not recipient:
            recipient = _recipient_from(address)
        numbers = extract_tracking_numbers(tracking_raw)
        if not any([tracking_raw, carrier_raw, address, notes]):
            continue
        skip_reason = None
        if not numbers:
            skip_reason = "missing_tracking"
            missing += 1
        elif is_delivered_text(status_raw):
            skip_reason = "already_delivered"
            delivered += 1
        else:
            pending += 1
        rows.append(
            SheetRow(
                excel_row=excel_row,
                recipient=recipient,
                carrier_raw=carrier_raw,
                tracking_raw=tracking_raw,
                tracking_numbers=numbers,
                status_raw=status_raw,
                address=address,
                notes=notes,
                skip_reason=skip_reason,
            )
        )
    return SheetTable(
        path=path,
        headers=headers,
        columns=columns,
        rows=rows,
        delivered_count=delivered,
        pending_count=pending,
        missing_count=missing,
    )


def apply_results(
    source: Path,
    dest: Path,
    updates: dict[int, dict[str, str]],
) -> Path:
    wb: Workbook = load_workbook(source)
    ws = wb.active
    columns = detect_columns(ws)
    if not _cell_text(ws.cell(1, columns.status).value):
        ws.cell(1, columns.status).value = "订单状态"
    if columns.latest > (ws.max_column or 1) or not _cell_text(ws.cell(1, columns.latest).value):
        ws.cell(1, columns.latest).value = "最新轨迹"
    if columns.queried_at > (ws.max_column or 1) or not _cell_text(ws.cell(1, columns.queried_at).value):
        ws.cell(1, columns.queried_at).value = "查询时间"

    now = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M")
    for excel_row, payload in updates.items():
        if payload.get("status"):
            ws.cell(excel_row, columns.status).value = payload["status"]
        if payload.get("latest"):
            ws.cell(excel_row, columns.latest).value = payload["latest"]
        ws.cell(excel_row, columns.queried_at).value = payload.get("queried_at") or now
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest
